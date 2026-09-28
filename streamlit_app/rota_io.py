"""Rota Excel load / normalize / title helpers."""

from __future__ import annotations

from datetime import datetime, time, timedelta

import numpy as np
import pandas as pd
import re
from openpyxl import load_workbook

from constants import UNIDADE_DISPLAY

ROTA_COLUMN_MAP = {
    "Grupo": "nome grupo",
    "Horário": "horario",
    "Unidade": "unidade",
    "STATUS": "status",
    "Dias da Semana": "dias da semana",
    "STAGE": "stage",
    "MODALIDADE": "modalidade",
    "GRUPO": "grupo",
    "N Aulas": "n aulas",
    "PARAG ATUAL": "parag atual grupo",
    "PARAG FINAL": "parag_final_grupo",
    "TEACHERS": "teacher",
    "Rescisão": "rescisao",
    "Permuta": "permuta",
    "Bolsista": "bolsista",
    "N TOTAL ALUNOS": "n_total_alunos",
}

ROTA_DISPLAY_MAP = {v: k for k, v in ROTA_COLUMN_MAP.items()}

ROTA_EXPORT_COLUMNS = [
    "Grupo",
    "Horário",
    "Unidade",
    "STATUS",
    "Dias da Semana",
    "STAGE",
    "MODALIDADE",
    "GRUPO",
    "N Aulas",
    "PARAG ATUAL",
    "PARAG FINAL",
    "TEACHERS",
    "N",
    "Rescisão",
    "Permuta",
    "Bolsista",
    "N TOTAL ALUNOS",
]


def normalize_horario_value(value):
    if pd.isna(value):
        return np.nan
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, datetime):
        return value.strftime("%H:%M:%S")
    if isinstance(value, pd.Timestamp):
        return value.strftime("%H:%M:%S")
    if isinstance(value, timedelta):
        total_seconds = int(value.total_seconds())
        hours, remainder = divmod(total_seconds, 3600)
        minutes, seconds = divmod(remainder, 60)
        return f"{hours:02d}:{minutes:02d}:{seconds:02d}"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        parsed = pd.to_datetime(value, unit="D", origin="1899-12-30", errors="coerce")
        if pd.notna(parsed):
            return parsed.strftime("%H:%M:%S")

    text = str(value).strip()
    if not text or text.lower() in ("nan", "none", "<na>"):
        return np.nan

    parsed = pd.to_datetime(text, errors="coerce")
    if pd.notna(parsed):
        return parsed.strftime("%H:%M:%S")
    return text


def normalize_stage(value):
    if pd.isna(value):
        return value
    if isinstance(value, (int, np.integer)):
        return f"ESTAGIO_{value}"
    if isinstance(value, float) and not np.isnan(value) and value == int(value):
        return f"ESTAGIO_{int(value)}"
    text = str(value).strip()
    if text.isdigit():
        return f"ESTAGIO_{text}"
    return text


def normalize_teacher_name(value) -> str | None:
    """Return a clean teacher name, or None when the cell means 'empty'."""
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    text = str(value).strip()
    if text.lower() in ("", "-", "nan", "none", "<na>", "nat"):
        return None
    return text


def teacher_or_dash(value) -> str:
    return normalize_teacher_name(value) or "-"


def stringify_column_name(column) -> str:
    if hasattr(column, "strftime"):
        return column.strftime("%H:%M:%S")
    return str(column).strip()


def normalize_rota_values(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()

    if "unidade" in df.columns:
        df["unidade"] = df["unidade"].astype(str).str.strip().str.upper()

    if "status" in df.columns:
        df["status"] = df["status"].astype(str).str.strip().str.upper()
        df["status"] = df["status"].replace({"NAN": np.nan, "NONE": np.nan})

    if "horario" in df.columns:
        df["horario"] = df["horario"].apply(normalize_horario_value)

    if "stage" in df.columns:
        df["stage"] = df["stage"].apply(normalize_stage)

    if "teacher" in df.columns:
        df["teacher"] = df["teacher"].map(teacher_or_dash)

    if "ultimo_professor" not in df.columns:
        df["ultimo_professor"] = "-"
    else:
        df["ultimo_professor"] = df["ultimo_professor"].map(teacher_or_dash)

    if "penultimo_professor" not in df.columns:
        df["penultimo_professor"] = "-"
    else:
        df["penultimo_professor"] = df["penultimo_professor"].map(teacher_or_dash)

    if "restricoes_professor" not in df.columns:
        df["restricoes_professor"] = np.nan

    df = df.loc[~df["nome grupo"].isnull()].copy()
    df = df[df["nome grupo"].astype(str).str.strip() != ""].copy()

    return df


def find_rota_sheet(wb) -> tuple[str, int, str | None]:
    if wb.active["A2"].value == "Grupo":
        a1 = wb.active["A1"].value
        titulo = str(a1).strip() if a1 and "ROTA" in str(a1).upper() else None
        return wb.active.title, 1, titulo

    for name in wb.sheetnames:
        ws = wb[name]
        if ws["A2"].value == "Grupo":
            a1 = ws["A1"].value
            titulo = str(a1).strip() if a1 and "ROTA" in str(a1).upper() else None
            return name, 1, titulo
        if ws["A1"].value == "Grupo":
            return name, 0, None

    raise ValueError(
        'Nenhuma aba de rota encontrada. Esperado cabeçalho "Grupo" na linha 1 ou 2.'
    )


HISTORY_TITLE_HEADERS = {
    "ultimo_professor",
    "penultimo_professor",
    "restricoes_professor",
}


def _read_title_row_history_headers(ws) -> dict[int, str]:
    """History headers sit on the title row (row 1), not the Grupo header row."""
    mapping: dict[int, str] = {}
    for row in ws.iter_rows(min_row=1, max_row=1):
        for cell in row:
            if cell.value is None:
                continue
            name = str(cell.value).strip()
            if name in HISTORY_TITLE_HEADERS:
                mapping[cell.column] = name
    return mapping


def load_rota_excel(file) -> tuple[pd.DataFrame, str | None, str]:
    """
    Lê o Excel da rota no formato do usuário (linha mesclada + cabeçalho)
    e normaliza colunas para o padrão interno do sistema.
    """
    wb = load_workbook(file, read_only=True, data_only=True)
    if hasattr(file, "seek"):
        file.seek(0)

    sheet_name, skiprows, titulo = find_rota_sheet(wb)
    history_headers = _read_title_row_history_headers(wb[sheet_name])
    wb.close()

    df = pd.read_excel(file, sheet_name=sheet_name, skiprows=skiprows)

    # Rename columns whose headers live on the title row (Excel 1-based index).
    rename_history = {}
    for excel_col, name in history_headers.items():
        pandas_idx = excel_col - 1
        if 0 <= pandas_idx < len(df.columns):
            rename_history[df.columns[pandas_idx]] = name
    if rename_history:
        df = df.rename(columns=rename_history)

    if "nome grupo" not in df.columns:
        df = df.rename(columns=ROTA_COLUMN_MAP)

    return normalize_rota_values(df), titulo, sheet_name


def parse_rota_titulo(titulo: str | None) -> tuple[str | None, str | None]:
    if not titulo:
        return None, None
    match = re.search(
        r"ROTA\s+(\d{2}/\d{2})\s+A\s+(\d{2}/\d{2})",
        str(titulo).strip(),
        re.IGNORECASE,
    )
    if match:
        return match.group(1), match.group(2)
    return None, None


def parse_dd_mm_to_date(dd_mm: str):
    day, month = dd_mm.split("/")
    return datetime(datetime.now().year, int(month), int(day)).date()


def build_rota_titulo(data_inicio: str, data_fim: str) -> str:
    inicio = data_inicio.strip()
    fim = data_fim.strip()
    if inicio and fim:
        return f"ROTA {inicio} A {fim}"
    if inicio:
        return f"ROTA {inicio}"
    return f"ROTA GERADA {datetime.now().strftime('%d/%m/%Y')}"


def format_rota_for_export(df: pd.DataFrame) -> pd.DataFrame:
    export_df = df.rename(columns=ROTA_DISPLAY_MAP)

    if "Unidade" in export_df.columns:
        export_df["Unidade"] = export_df["Unidade"].map(
            lambda value: UNIDADE_DISPLAY.get(str(value).strip().upper(), value)
        )

    if "TEACHERS" in export_df.columns:
        export_df["TEACHERS"] = export_df["TEACHERS"].replace("-", "")

    cols = [col for col in ROTA_EXPORT_COLUMNS if col in export_df.columns]
    extra_cols = [col for col in export_df.columns if col not in cols]
    return export_df[cols + extra_cols]


def teacher_display_value(teacher) -> str | None:
    if teacher is None or pd.isna(teacher):
        return None
    value = str(teacher).strip()
    if value in ("", "-", "nan", "NaN", "None", "<NA>"):
        return None
    return value
