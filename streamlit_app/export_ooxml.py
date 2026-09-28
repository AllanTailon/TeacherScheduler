"""OOXML patching and Excel export for rota workbooks."""

from __future__ import annotations

import io
import re
import unicodedata
import zipfile
from datetime import datetime
from xml.etree import ElementTree as ET

import pandas as pd
from openpyxl.styles import Font, PatternFill

from rota_io import format_rota_for_export, teacher_display_value

SHEET_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
SHEET_NS = {"m": SHEET_MAIN_NS}
REL_ID_ATTR = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
ET.register_namespace("", SHEET_MAIN_NS)


def _resolve_workbook_sheet_path(zf: zipfile.ZipFile, sheet_name: str) -> str:
    wb = ET.fromstring(zf.read("xl/workbook.xml"))
    rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    rid_map = {rel.attrib["Id"]: rel.attrib["Target"] for rel in rels}
    for sheet in wb.findall(".//m:sheets/m:sheet", SHEET_NS):
        if sheet.attrib.get("name") == sheet_name:
            target = rid_map[sheet.attrib[REL_ID_ATTR]]
            return "xl/" + target.lstrip("/")
    raise ValueError(f'Aba "{sheet_name}" não encontrada no template.')


def _read_shared_strings(zf: zipfile.ZipFile) -> list[str]:
    if "xl/sharedStrings.xml" not in zf.namelist():
        return []
    root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
    strings = []
    for item in root.findall("m:si", SHEET_NS):
        text = item.find("m:t", SHEET_NS)
        if text is not None:
            strings.append(text.text or "")
            continue
        parts = []
        for run in item.findall("m:r", SHEET_NS):
            run_text = run.find("m:t", SHEET_NS)
            if run_text is not None and run_text.text:
                parts.append(run_text.text)
        strings.append("".join(parts))
    return strings


def _xml_escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _replace_si_text_at_index(shared_xml: bytes, index: int, new_text: str) -> bytes:
    text = shared_xml.decode("utf-8")
    matches = list(re.finditer(r"<si>(?:.*?)</si>", text, flags=re.DOTALL))
    if index >= len(matches):
        return shared_xml
    match = matches[index]
    new_si = f"<si><t>{_xml_escape(new_text)}</t></si>"
    return (text[: match.start()] + new_si + text[match.end() :]).encode("utf-8")


def _get_shared_string_index(cell) -> int | None:
    if cell is None or cell.attrib.get("t") != "s":
        return None
    value = cell.find("m:v", SHEET_NS)
    if value is None or value.text is None:
        return None
    return int(value.text)


def _normalize_grupo_key(value) -> str:
    return unicodedata.normalize("NFC", str(value).strip())


def _append_shared_string(shared_xml: bytes, new_text: str) -> tuple[bytes, int]:
    text = shared_xml.decode("utf-8")
    matches = list(re.finditer(r"<si>(?:.*?)</si>", text, flags=re.DOTALL))
    new_index = len(matches)
    new_si = f"<si><t>{_xml_escape(new_text)}</t></si>"
    text = text.replace("</sst>", new_si + "</sst>", 1)
    text = re.sub(
        r'uniqueCount="(\d+)"',
        lambda match: f'uniqueCount="{int(match.group(1)) + 1}"',
        text,
        count=1,
    )
    text = re.sub(
        r'count="(\d+)"',
        lambda match: f'count="{int(match.group(1)) + 1}"',
        text,
        count=1,
    )
    return text.encode("utf-8"), new_index


def _set_cell_shared_index(sheet_xml: bytes, cell_ref: str, index: int) -> bytes:
    text = sheet_xml.decode("utf-8")
    pattern = rf'<c r="{re.escape(cell_ref)}"([^>/]*)(?:/>|>(?:.*?</c>))'

    def replacer(match):
        attrs = match.group(1)
        attrs = re.sub(r'\s*t="[^"]*"', "", attrs)
        return f'<c r="{cell_ref}"{attrs} t="s"><v>{index}</v></c>'

    new_text, replacements = re.subn(pattern, replacer, text, count=1, flags=re.DOTALL)
    if replacements:
        return new_text.encode("utf-8")

    row_number = int("".join(ch for ch in cell_ref if ch.isdigit()))
    row_pattern = rf'(<row r="{row_number}"[^>]*>)(.*?)(</row>)'
    row_match = re.search(row_pattern, text, flags=re.DOTALL)
    if row_match is None:
        return sheet_xml

    new_cell = f'<c r="{cell_ref}" t="s"><v>{index}</v></c>'
    updated_row = row_match.group(1) + row_match.group(2) + new_cell + row_match.group(3)
    return (text[: row_match.start()] + updated_row + text[row_match.end() :]).encode("utf-8")


def _cell_ref_parts(ref: str) -> tuple[str, int]:
    col = "".join(ch for ch in ref if ch.isalpha())
    row = int("".join(ch for ch in ref if ch.isdigit()))
    return col, row


def _find_row(sheet_root, row_number: int):
    for row in sheet_root.findall(".//m:sheetData/m:row", SHEET_NS):
        if int(row.attrib.get("r", 0)) == row_number:
            return row
    return None


def _find_cell(row, ref: str):
    for cell in row.findall("m:c", SHEET_NS):
        if cell.attrib.get("r") == ref:
            return cell
    return None


def _resolve_cell_text(cell, strings: list[str]) -> str | None:
    if cell is None:
        return None
    cell_type = cell.attrib.get("t")
    value = cell.find("m:v", SHEET_NS)
    if cell_type == "s" and value is not None and value.text is not None:
        return strings[int(value.text)]
    if cell_type == "inlineStr":
        inline = cell.find("m:is/m:t", SHEET_NS)
        return inline.text if inline is not None else None
    if value is not None:
        return value.text
    return None


def _find_header_column(
    sheet_root, strings: list[str], header_name: str, header_row: int = 2
) -> str | None:
    row = _find_row(sheet_root, header_row)
    if row is None:
        return None
    for cell in row.findall("m:c", SHEET_NS):
        if _resolve_cell_text(cell, strings) == header_name:
            return _cell_ref_parts(cell.attrib["r"])[0]
    return None


def _find_header_column_any_row(
    sheet_root, strings: list[str], header_name: str, rows=(1, 2)
) -> str | None:
    for header_row in rows:
        found = _find_header_column(sheet_root, strings, header_name, header_row)
        if found:
            return found
    return None


def _build_rota_patch_targets(
    sheet_xml: bytes,
    strings: list[str],
) -> tuple[int | None, dict[str, dict[str, str]]]:
    """Return title shared-string index and per-group cell refs for patch fields."""
    root = ET.fromstring(sheet_xml)
    header_row = 2 if _find_row(root, 2) is not None else 1
    data_start = header_row + 1
    grupo_col = _find_header_column(root, strings, "Grupo", header_row) or "A"
    teachers_col = _find_header_column(root, strings, "TEACHERS", header_row) or "L"
    ultimo_col = _find_header_column_any_row(root, strings, "ultimo_professor")
    penultimo_col = _find_header_column_any_row(root, strings, "penultimo_professor")

    title_row = _find_row(root, 1)
    title_ss_index = (
        _get_shared_string_index(_find_cell(title_row, "A1")) if title_row is not None else None
    )

    cells_by_grupo: dict[str, dict[str, str]] = {}
    for row in root.findall(".//m:sheetData/m:row", SHEET_NS):
        row_number = int(row.attrib.get("r", 0))
        if row_number < data_start:
            continue
        grupo = _resolve_cell_text(_find_cell(row, f"{grupo_col}{row_number}"), strings)
        if not grupo:
            continue
        key = _normalize_grupo_key(grupo)
        refs = {"teacher": f"{teachers_col}{row_number}"}
        if ultimo_col:
            refs["ultimo"] = f"{ultimo_col}{row_number}"
        if penultimo_col:
            refs["penultimo"] = f"{penultimo_col}{row_number}"
        cells_by_grupo[key] = refs

    return title_ss_index, cells_by_grupo


def _patch_rota_workbook(
    shared_xml: bytes,
    sheet_xml: bytes,
    titulo: str,
    values_by_grupo: dict[str, dict[str, str | None]],
    title_ss_index: int | None,
    cells_by_grupo: dict[str, dict[str, str]],
) -> tuple[bytes, bytes]:
    patched_shared = shared_xml
    patched_sheet = sheet_xml
    string_cache: dict[str, int] = {}

    if title_ss_index is not None:
        patched_shared = _replace_si_text_at_index(patched_shared, title_ss_index, titulo)

    def _write_value(cell_ref: str, text_value: str) -> None:
        nonlocal patched_shared, patched_sheet
        if text_value not in string_cache:
            patched_shared, string_cache[text_value] = _append_shared_string(
                patched_shared, text_value
            )
        patched_sheet = _set_cell_shared_index(
            patched_sheet, cell_ref, string_cache[text_value]
        )

    for grupo, values in values_by_grupo.items():
        refs = cells_by_grupo.get(_normalize_grupo_key(grupo))
        if not refs:
            continue
        teacher = values.get("teacher")
        if teacher and "teacher" in refs:
            _write_value(refs["teacher"], teacher)
        ultimo = values.get("ultimo")
        if ultimo and "ultimo" in refs:
            _write_value(refs["ultimo"], ultimo)
        penultimo = values.get("penultimo")
        if penultimo and "penultimo" in refs:
            _write_value(refs["penultimo"], penultimo)

    return patched_shared, patched_sheet


def _clone_zip_with_patches(template_bytes: bytes, patches: dict[str, bytes]) -> bytes:
    in_zip = zipfile.ZipFile(io.BytesIO(template_bytes), "r")
    out_buf = io.BytesIO()
    with zipfile.ZipFile(out_buf, "w") as out_zip:
        for item in in_zip.infolist():
            data = patches.get(item.filename, in_zip.read(item.filename))
            out_zip.writestr(item, data, compress_type=item.compress_type)
    in_zip.close()
    return out_buf.getvalue()


def _export_rota_from_template(
    df: pd.DataFrame,
    titulo: str,
    template_bytes: bytes,
    sheet_name: str | None = None,
) -> bytes:
    values_by_grupo: dict[str, dict[str, str | None]] = {}
    for _, row in df.iterrows():
        if pd.isna(row.get("nome grupo")):
            continue
        grupo = _normalize_grupo_key(row["nome grupo"])
        values_by_grupo[grupo] = {
            "teacher": teacher_display_value(row.get("teacher")),
            "ultimo": teacher_display_value(row.get("ultimo_professor")),
            "penultimo": teacher_display_value(row.get("penultimo_professor")),
        }

    with zipfile.ZipFile(io.BytesIO(template_bytes), "r") as zf:
        if sheet_name is None:
            wb = ET.fromstring(zf.read("xl/workbook.xml"))
            active_id = wb.find(".//m:bookViews/m:workbookView", SHEET_NS).attrib.get(
                "activeTab", "0"
            )
            sheets = wb.findall(".//m:sheets/m:sheet", SHEET_NS)
            sheet_name = sheets[int(active_id)].attrib["name"]
        sheet_path = _resolve_workbook_sheet_path(zf, sheet_name)
        strings = _read_shared_strings(zf)
        sheet_xml = zf.read(sheet_path)
        shared_xml = zf.read("xl/sharedStrings.xml")

    title_ss_index, cells_by_grupo = _build_rota_patch_targets(sheet_xml, strings)
    patched_shared, patched_sheet = _patch_rota_workbook(
        shared_xml,
        sheet_xml,
        titulo,
        values_by_grupo,
        title_ss_index,
        cells_by_grupo,
    )
    return _clone_zip_with_patches(
        template_bytes,
        {
            "xl/sharedStrings.xml": patched_shared,
            sheet_path: patched_sheet,
        },
    )


def _export_rota_plain(df: pd.DataFrame, titulo: str) -> bytes:
    export_df = format_rota_for_export(df)

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame([[titulo]]).to_excel(
            writer, index=False, header=False, startrow=0, sheet_name="Rotas"
        )
        export_df.to_excel(writer, index=False, sheet_name="Rotas", startrow=1)
        ws = writer.sheets["Rotas"]
        ws.merge_cells(
            start_row=1, start_column=1, end_row=1, end_column=max(len(export_df.columns), 1)
        )

        header_fill = PatternFill(patternType="solid", fgColor="FF000000")
        header_font = Font(color="FFFFFFFF", bold=True)
        title_font = Font(color="FFFFFFFF", bold=True)
        data_fill = PatternFill(patternType="solid", fgColor="FFFFFFFF")
        green_header_fill = PatternFill(patternType="solid", fgColor="FF8BC34A")
        green_header_font = Font(color="FF000000", bold=True)

        for col in range(1, len(export_df.columns) + 1):
            cell = ws.cell(row=2, column=col)
            if col == len(export_df.columns) and export_df.columns[-1] == "N TOTAL ALUNOS":
                cell.fill = green_header_fill
                cell.font = green_header_font
            else:
                cell.fill = header_fill
                cell.font = header_font

        ws["A1"].fill = header_fill
        ws["A1"].font = title_font

        for row_idx in range(3, 3 + len(export_df)):
            for col in range(1, len(export_df.columns) + 1):
                ws.cell(row=row_idx, column=col).fill = data_fill

    return output.getvalue()


def export_rota_excel(
    df: pd.DataFrame,
    titulo: str | None = None,
    template_bytes: bytes | None = None,
    sheet_name: str | None = None,
) -> bytes:
    if not titulo:
        titulo = f"ROTA GERADA {datetime.now().strftime('%d/%m/%Y')}"

    if template_bytes:
        return _export_rota_from_template(df, titulo, template_bytes, sheet_name)
    return _export_rota_plain(df, titulo)


def _set_row_hidden_flag(sheet_text: str, row_num: int, hidden: bool) -> str:
    pattern = rf'(<row r="{row_num}"[^>]*?)>'
    match = re.search(pattern, sheet_text)
    if not match:
        return sheet_text
    row_tag = re.sub(r'\s*hidden="1"', "", match.group(1))
    if hidden:
        row_tag += ' hidden="1"'
    return sheet_text[: match.start()] + row_tag + ">" + sheet_text[match.end() :]


def _filter_rota_sheet_for_groups(
    sheet_xml: bytes,
    strings: list[str],
    allowed_grupos: set[str],
) -> bytes:
    _, cells_by_grupo = _build_rota_patch_targets(sheet_xml, strings)
    allowed = {_normalize_grupo_key(grupo) for grupo in allowed_grupos}
    visible_rows = {1, 2}

    for grupo, refs in cells_by_grupo.items():
        if grupo in allowed:
            cell_ref = refs.get("teacher")
            if cell_ref:
                visible_rows.add(int("".join(ch for ch in cell_ref if ch.isdigit())))

    sheet_text = sheet_xml.decode("utf-8")
    for row_match in re.finditer(r'<row r="(\d+)"', sheet_text):
        row_num = int(row_match.group(1))
        sheet_text = _set_row_hidden_flag(sheet_text, row_num, row_num not in visible_rows)

    return sheet_text.encode("utf-8")


def export_rota_for_teacher(
    template_bytes: bytes,
    allowed_grupos: list[str],
    sheet_name: str | None = None,
) -> bytes:
    with zipfile.ZipFile(io.BytesIO(template_bytes), "r") as zf:
        if sheet_name is None:
            wb = ET.fromstring(zf.read("xl/workbook.xml"))
            active_id = wb.find(".//m:bookViews/m:workbookView", SHEET_NS).attrib.get(
                "activeTab", "0"
            )
            sheets = wb.findall(".//m:sheets/m:sheet", SHEET_NS)
            sheet_name = sheets[int(active_id)].attrib["name"]
        sheet_path = _resolve_workbook_sheet_path(zf, sheet_name)
        strings = _read_shared_strings(zf)
        sheet_xml = zf.read(sheet_path)

    filtered_sheet = _filter_rota_sheet_for_groups(sheet_xml, strings, set(allowed_grupos))
    return _clone_zip_with_patches(template_bytes, {sheet_path: filtered_sheet})
